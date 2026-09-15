#!/usr/bin/env python3
"""Шаблон: генерация XLSX тест-кейсов для импорта в Test IT (15 колонок).

Скопируй файл, замени данные в CASES на свои и запусти:

    python3 test_it_generator.py /путь/к/output.xlsx
"""

import os
import sys
import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---------------------------------------------------------------------------
# Данные тест-кейсов — ЗАМЕНИ НА СВОИ
# ---------------------------------------------------------------------------
BASE_SRC = "Путь_к_требованию"

CASES = [
    {
        "name": "TC-001 – Название тест-кейса",
        "desc": f"Источник: {BASE_SRC}\nТребование: Описание требования",
        "priority": "Высокий",
        "duration": "0h 5m 0s",
        "section": "Секция > Подсекция",
        "review": "Не пройдено",
        "type": "Позитивный",
        "preconditions": [
            "Предусловие 1",
            "Предусловие 2",
        ],
        "steps": [
            {
                "action": "Шаг 1 — действие пользователя",
                "expected": "Ожидаемый результат 1",
            },
            {
                "action": "Шаг 2 — действие пользователя",
                "expected": "Ожидаемый результат 2",
            },
        ],
        "postconditions": [],
    },
    # Пункт чек-листа: без preconditions и steps.
    # Даёт одну Header-строку, колонки «Действие» и «Ожидаемый результат» пустые.
    {
        "name": "Успешная регистрация пользователя",
        "priority": "Высокий",
        "duration": "0h 3m 0s",
        "section": "Секция > Подсекция",
        "review": "Не пройдено",
        "type": "Позитивный",
    },
]

DEFAULT_OUTPUT = "TC_Output.xlsx"

# ---------------------------------------------------------------------------
# Структура и стилизация
# ---------------------------------------------------------------------------
COLUMNS = [
    "Название",                          # 1
    "Описание",                          # 2
    "Статус",                            # 3
    "Приоритет",                         # 4
    "Действие",                          # 5
    "Ожидаемый результат",               # 6
    "Действие предусловия",              # 7
    "Ожидаемый результат предусловия",   # 8
    "Действие постусловия",              # 9
    "Ожидаемый результат постусловия",   # 10
    "Продолжительность",                 # 11
    "Секция",                            # 12
    "Тестовые данные",                   # 13
    "Ревью",                             # 14
    "Вид",                               # 15
]

COL_WIDTHS = [50, 60, 10, 12, 55, 65, 55, 10, 55, 10, 14, 30, 14, 12, 12]

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
TITLE_FONT = Font(name="Calibri", bold=True, size=11)
BODY_FONT = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def write_row(ws, row: int, values: list, font: Font, fill: PatternFill = None):
    """Записывает строку из 15 значений с единым оформлением."""
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.font = font
        cell.alignment = WRAP
        cell.border = THIN_BORDER
        if fill is not None:
            cell.fill = fill


def blank_row() -> list:
    return [""] * len(COLUMNS)


def write_case(ws, start_row: int, case: dict) -> int:
    """Пишет один тест-кейс. Возвращает номер следующей свободной строки."""
    row = start_row

    # Header-строка: заполнены только обязательные поля
    header = blank_row()
    header[0] = case["name"]              # Название
    header[1] = case.get("desc", "")      # Описание
    header[2] = "Готов"             # Статус
    header[3] = case["priority"]    # Приоритет
    header[10] = case["duration"]   # Продолжительность
    header[11] = case["section"]    # Секция
    header[13] = case["review"]     # Ревью
    header[14] = case["type"]       # Вид
    write_row(ws, row, header, HEADER_FONT, HEADER_FILL)
    row += 1

    # Предусловия
    for precondition in case.get("preconditions", []):
        values = blank_row()
        values[6] = precondition
        write_row(ws, row, values, BODY_FONT)
        row += 1

    # Шаги и ожидаемые результаты.
    # У пункта чек-листа ожидаемого результата нет — колонка остаётся пустой.
    for step in case.get("steps", []):
        values = blank_row()
        values[4] = step["action"]
        values[5] = step.get("expected", "")
        write_row(ws, row, values, BODY_FONT)
        row += 1

    # Постусловия
    for postcondition in case.get("postconditions", []):
        values = blank_row()
        values[8] = postcondition["action"]
        values[9] = postcondition["expected"]
        write_row(ws, row, values, BODY_FONT)
        row += 1

    # Разделитель нужен только когда у Header-строки есть дочерние строки.
    # Пункт чек-листа — одна строка, разделители между такими строками не ставятся.
    return row + 1 if row > start_row + 1 else row


def main():
    output_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_OUTPUT

    wb = Workbook()
    ws = wb.active
    ws.title = "Тест-кейсы"

    write_row(ws, 1, COLUMNS, TITLE_FONT, TITLE_FILL)

    row = 2
    for case in CASES:
        row = write_case(ws, row, case)

    for col_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    ws.freeze_panes = "A2"

    wb.save(output_path)
    time.sleep(1)

    exists = os.path.exists(output_path)
    size = os.path.getsize(output_path) if exists else 0
    print(f"[{output_path}] exists={exists} size={size} bytes cases={len(CASES)}")


if __name__ == "__main__":
    main()
